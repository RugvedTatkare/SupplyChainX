"""Safe local storage and normalization for user supplied supply-chain datasets."""
from __future__ import annotations

import csv
import io
import json
import math
import re
import sqlite3
import uuid
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_PREVIEW_ROWS = 12
MAX_XLSX_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
MAX_ZIP_RATIO = 200
DATASETS_DIR = Path(__file__).resolve().parents[2] / "data" / "imported_datasets"

# Fields that are safe and useful to display or aggregate. Risk scores are
# deliberately excluded: imports never create predictive or prototype scores.
FIELD_LABELS = {
    "supplier_id": "Supplier ID",
    "supplier_name": "Supplier name",
    "supplier_country": "Supplier country",
    "supplier_late_delivery_rate": "Supplier late-delivery rate (%)",
    "shipment_id": "Shipment ID",
    "order_id": "Order ID",
    "ship_date": "Ship date",
    "scheduled_delivery_date": "Scheduled delivery date",
    "actual_delivery_date": "Actual delivery date",
    "scheduled_shipping_days": "Scheduled shipping days",
    "actual_shipping_days": "Actual shipping days",
    "delivery_delay_days": "Delivery delay (days)",
    "delivery_status": "Delivery status",
    "shipping_mode": "Shipping mode",
    "warehouse_id": "Warehouse ID",
    "order_country": "Order country",
    "order_region": "Order region",
    "product_id": "Product ID",
    "product_name": "Product name",
    "snapshot_date": "Inventory snapshot date",
    "inventory_quantity": "Inventory quantity",
    "reorder_point": "Reorder point",
    "daily_demand": "Daily demand",
}

IDENTITY_FIELDS = {"supplier_id", "supplier_name", "shipment_id", "order_id", "product_id", "warehouse_id"}
MEASURE_FIELDS = {
    "supplier_late_delivery_rate", "scheduled_delivery_date", "actual_delivery_date",
    "scheduled_shipping_days", "actual_shipping_days", "delivery_delay_days",
    "delivery_status", "inventory_quantity", "reorder_point", "daily_demand",
}
NUMERIC_FIELDS = {
    "supplier_late_delivery_rate", "scheduled_shipping_days", "actual_shipping_days",
    "delivery_delay_days", "inventory_quantity", "reorder_point", "daily_demand",
}
DATE_FIELDS = {"ship_date", "scheduled_delivery_date", "actual_delivery_date", "snapshot_date"}

ALIASES = {
    "supplier_id": {"supplierid", "vendorid", "suppliercode", "vendorcode"},
    "supplier_name": {"suppliername", "vendorname", "supplier", "vendor"},
    "supplier_country": {"suppliercountry", "vendorcountry", "countryofsupplier"},
    "supplier_late_delivery_rate": {"supplierlatedeliveryrate", "latedeliveryrate", "late delivery rate", "supplierdelayrate"},
    "shipment_id": {"shipmentid", "trackingid", "shipmentnumber", "trackingnumber"},
    "order_id": {"orderid", "ordernumber", "purchaseorder"},
    "ship_date": {"shipdate", "shipmentdate", "shippeddate", "shipdatetime"},
    "scheduled_delivery_date": {"scheduleddeliverydate", "expecteddeliverydate", "promiseddeliverydate"},
    "actual_delivery_date": {"actualdeliverydate", "deliverydate", "delivereddate"},
    "scheduled_shipping_days": {"scheduledshippingdays", "plannedshippingdays", "scheduleddays"},
    "actual_shipping_days": {"actualshippingdays", "shippingdays", "transitdays"},
    "delivery_delay_days": {"deliverydelaydays", "delaydays", "dayslate", "deliverydelay"},
    "delivery_status": {"deliverystatus", "status", "shipmentstatus", "fulfillmentstatus"},
    "shipping_mode": {"shippingmode", "transportmode", "carrier"},
    "warehouse_id": {"warehouseid", "warehouse", "facilityid"},
    "order_country": {"ordercountry", "destinationcountry", "country"},
    "order_region": {"orderregion", "region", "destinationregion"},
    "product_id": {"productid", "sku", "itemid", "productcode"},
    "product_name": {"productname", "itemname", "product"},
    "snapshot_date": {"snapshotdate", "inventorydate", "asofdate"},
    "inventory_quantity": {"inventoryquantity", "currentstock", "stockquantity", "quantityonhand", "onhand", "stock", "inventory"},
    "reorder_point": {"reorderpoint", "reorderlevel", "minimumstock"},
    "daily_demand": {"dailydemand", "estimateddailydemand", "unitsperday"},
}


class DatasetValidationError(ValueError):
    """A safe, user-displayable validation failure."""


def _clean_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _as_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def parse_upload(filename: str | None, content: bytes) -> tuple[list[str], list[list[str]]]:
    """Decode bounded CSV/XLSX content into a header and rows."""
    safe_name = Path(filename or "").name
    suffix = Path(safe_name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise DatasetValidationError("Choose a CSV or Excel (.xlsx) file.")
    if not content:
        raise DatasetValidationError("The selected file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise DatasetValidationError("The file exceeds the 20 MB upload limit.")

    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DatasetValidationError("CSV files must use UTF-8 encoding.") from exc
        if "\x00" in text:
            raise DatasetValidationError("The CSV contains unsupported binary characters.")
        try:
            reader = csv.reader(io.StringIO(text, newline=""), strict=True)
            raw_rows = list(reader)
        except csv.Error as exc:
            raise DatasetValidationError(f"Malformed CSV: {exc}") from exc
    else:
        if not content.startswith(b"PK\x03\x04"):
            raise DatasetValidationError("The file extension is .xlsx, but the content is not an Excel workbook.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                expanded = sum(entry.file_size for entry in entries)
                if expanded > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise DatasetValidationError("The Excel workbook expands beyond the safe processing limit.")
                if any(entry.compress_size and entry.file_size / entry.compress_size > MAX_ZIP_RATIO for entry in entries):
                    raise DatasetValidationError("The Excel workbook has an unsafe compression ratio.")
                if "xl/workbook.xml" not in archive.namelist():
                    raise DatasetValidationError("The selected .xlsx file is not a valid Excel workbook.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                sheet = next((item for item in workbook.worksheets if item.sheet_state == "visible"), None)
                if sheet is None:
                    raise DatasetValidationError("The workbook has no visible worksheet to import.")
                if sheet.max_row > MAX_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
                    raise DatasetValidationError(f"Excel sheets are limited to {MAX_ROWS:,} rows and {MAX_COLUMNS} columns.")
                raw_rows = [[_as_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            finally:
                workbook.close()
        except DatasetValidationError:
            raise
        except Exception as exc:
            raise DatasetValidationError("The selected .xlsx file could not be read as a valid workbook.") from exc

    while raw_rows and not any(_clean_header(cell) for cell in raw_rows[0]):
        raw_rows.pop(0)
    if not raw_rows:
        raise DatasetValidationError("The file has no header row.")
    headers = [_clean_header(cell) for cell in raw_rows[0]]
    if len(headers) > MAX_COLUMNS:
        raise DatasetValidationError(f"Files are limited to {MAX_COLUMNS} columns.")
    if any(not header for header in headers):
        raise DatasetValidationError("Every column needs a name. Rename blank columns and try again.")
    normalized = [header.casefold() for header in headers]
    if len(normalized) != len(set(normalized)):
        raise DatasetValidationError("Column names must be unique, ignoring capitalization.")
    rows: list[list[str]] = []
    for line_number, row in enumerate(raw_rows[1:], start=2):
        if not any(_as_cell(value) for value in row):
            continue
        if len(row) != len(headers):
            raise DatasetValidationError(f"Row {line_number} has {len(row)} values; expected {len(headers)}.")
        rows.append([_as_cell(value) for value in row])
        if len(rows) > MAX_ROWS:
            raise DatasetValidationError(f"Files are limited to {MAX_ROWS:,} non-empty rows.")
    if not rows:
        raise DatasetValidationError("The file has column headers but no data rows.")
    return headers, rows


def suggest_mappings(headers: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    claimed: set[str] = set()
    for field, aliases in ALIASES.items():
        candidates = {re.sub(r"[^a-z0-9]", "", alias.casefold()) for alias in aliases | {field}}
        for header in headers:
            normalized = re.sub(r"[^a-z0-9]", "", header.casefold())
            if normalized in candidates and header not in claimed:
                result[field] = header
                claimed.add(header)
                break
    return result


def _parse_date(value: str, field: str, row_number: int) -> str:
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return parsed.date().isoformat()
    except ValueError:
        pass
    for pattern in ("%m/%d/%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, pattern).date().isoformat()
        except ValueError:
            continue
    raise DatasetValidationError(f"Row {row_number}: {FIELD_LABELS[field]} must be a valid date (for example, 2026-03-21).")


def normalize_rows(
    headers: list[str], rows: list[list[str]], mappings: dict[str, Any],
) -> tuple[dict[str, str], list[dict[str, Any]]]:
    if not isinstance(mappings, dict):
        raise DatasetValidationError("Column mappings must be an object.")
    unknown_fields = sorted(set(mappings) - set(FIELD_LABELS))
    if unknown_fields:
        raise DatasetValidationError(f"Unsupported mapped fields: {', '.join(unknown_fields)}.")
    selected: dict[str, str] = {}
    used_headers: dict[str, str] = {}
    by_case = {header.casefold(): header for header in headers}
    for field, raw_header in mappings.items():
        if raw_header in (None, ""):
            continue
        if not isinstance(raw_header, str) or raw_header.casefold() not in by_case:
            raise DatasetValidationError(f"The selected source column for {FIELD_LABELS[field]} does not exist in this file.")
        header = by_case[raw_header.casefold()]
        if header in used_headers:
            raise DatasetValidationError(f"{header} is assigned to both {FIELD_LABELS[used_headers[header]]} and {FIELD_LABELS[field]}.")
        selected[field] = header
        used_headers[header] = field
    if not IDENTITY_FIELDS.intersection(selected):
        raise DatasetValidationError("Map at least one identifier: supplier, shipment, order, product, or warehouse.")
    if not MEASURE_FIELDS.intersection(selected):
        raise DatasetValidationError("Map at least one measurable field: delivery status or dates, delay, inventory quantity, demand, or reorder point.")

    positions = {field: headers.index(header) for field, header in selected.items()}
    normalized_rows: list[dict[str, Any]] = []
    identity_fields = IDENTITY_FIELDS.intersection(selected)
    measure_fields = MEASURE_FIELDS.intersection(selected)
    for row_number, row in enumerate(rows, start=2):
        output: dict[str, Any] = {}
        for field, position in positions.items():
            value = row[position].strip()
            if not value:
                output[field] = None
                continue
            if len(value) > 2_000:
                raise DatasetValidationError(f"Row {row_number}: values are limited to 2,000 characters.")
            if field in NUMERIC_FIELDS:
                try:
                    numeric = float(value.replace(",", ""))
                except ValueError as exc:
                    raise DatasetValidationError(f"Row {row_number}: {FIELD_LABELS[field]} must be numeric.") from exc
                if not math.isfinite(numeric):
                    raise DatasetValidationError(f"Row {row_number}: {FIELD_LABELS[field]} must be a finite number.")
                if field == "supplier_late_delivery_rate" and not 0 <= numeric <= 100:
                    raise DatasetValidationError(f"Row {row_number}: {FIELD_LABELS[field]} must be between 0 and 100.")
                output[field] = numeric
            elif field in DATE_FIELDS:
                output[field] = _parse_date(value, field, row_number)
            else:
                output[field] = value
        if not any(output.get(field) for field in identity_fields):
            raise DatasetValidationError(f"Row {row_number}: a supplier, shipment, order, product, or warehouse identifier is required.")
        if not any(output.get(field) is not None for field in measure_fields):
            raise DatasetValidationError(f"Row {row_number}: at least one mapped measurement is required.")
        normalized_rows.append(output)
    return selected, normalized_rows


def _safe_filename(filename: str | None) -> str:
    name = Path(filename or "dataset").name
    name = re.sub(r"[^A-Za-z0-9._ -]", "", name).strip(" .")
    return (name or "dataset")[:120]


def import_dataset(filename: str | None, headers: list[str], rows: list[list[str]], mappings: dict[str, Any]) -> dict[str, Any]:
    selected, normalized_rows = normalize_rows(headers, rows, mappings)
    dataset_id = uuid.uuid4().hex
    safe_name = _safe_filename(filename)
    target_dir = DATASETS_DIR
    target_dir.mkdir(parents=True, exist_ok=True)
    temporary = target_dir / f".{dataset_id}.tmp"
    destination = target_dir / f"{dataset_id}.sqlite"
    metadata = {
        "dataset_id": dataset_id,
        "name": Path(safe_name).stem or "Imported dataset",
        "filename": safe_name,
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "row_count": len(normalized_rows),
        "columns": headers,
        "mappings": selected,
    }
    try:
        connection = sqlite3.connect(temporary)
        try:
            connection.execute("PRAGMA journal_mode=DELETE")
            connection.execute("CREATE TABLE metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)")
            connection.execute("CREATE TABLE records (row_id INTEGER PRIMARY KEY, payload TEXT NOT NULL)")
            connection.executemany("INSERT INTO metadata (key, value) VALUES (?, ?)", [
                (key, json.dumps(value, ensure_ascii=False)) for key, value in metadata.items()
            ])
            connection.executemany("INSERT INTO records (payload) VALUES (?)", [
                (json.dumps(row, ensure_ascii=False, allow_nan=False),) for row in normalized_rows
            ])
            connection.commit()
        finally:
            connection.close()
        temporary.replace(destination)
    except Exception:
        temporary.unlink(missing_ok=True)
        destination.unlink(missing_ok=True)
        raise
    return dataset_summary(metadata)


def _dataset_path(dataset_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", dataset_id):
        raise DatasetValidationError("Unknown dataset selection.")
    path = (DATASETS_DIR / f"{dataset_id}.sqlite").resolve()
    if path.parent != DATASETS_DIR.resolve() or not path.is_file():
        raise DatasetValidationError("This imported dataset could not be found. It may have been removed locally.")
    return path


def _read_metadata(connection: sqlite3.Connection) -> dict[str, Any]:
    return {key: json.loads(value) for key, value in connection.execute("SELECT key, value FROM metadata")}


def dataset_summary(metadata: dict[str, Any]) -> dict[str, Any]:
    return {key: metadata[key] for key in ("dataset_id", "name", "filename", "created_at", "row_count", "mappings")}


def list_datasets() -> list[dict[str, Any]]:
    DATASETS_DIR.mkdir(parents=True, exist_ok=True)
    datasets: list[dict[str, Any]] = [{
        "dataset_id": "original", "name": "Original historical dataset", "filename": None,
        "created_at": None, "row_count": None, "mappings": {}, "kind": "original",
    }]
    for path in sorted(DATASETS_DIR.glob("[a-f0-9]" * 32 + ".sqlite"), key=lambda item: item.stat().st_mtime, reverse=True):
        try:
            connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
            try:
                metadata = _read_metadata(connection)
            finally:
                connection.close()
            datasets.append({**dataset_summary(metadata), "kind": "imported"})
        except (sqlite3.Error, KeyError, json.JSONDecodeError):
            continue
    return datasets


def read_imported_dataset(dataset_id: str) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    path = _dataset_path(dataset_id)
    connection = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
    try:
        metadata = _read_metadata(connection)
        rows = [json.loads(payload) for (payload,) in connection.execute("SELECT payload FROM records ORDER BY row_id")]
    finally:
        connection.close()
    return metadata, rows


def delete_dataset(dataset_id: str) -> None:
    path = _dataset_path(dataset_id)
    path.unlink()


def _outcome(row: dict[str, Any]) -> bool | None:
    if row.get("delivery_delay_days") is not None:
        return row["delivery_delay_days"] <= 0
    if row.get("scheduled_delivery_date") and row.get("actual_delivery_date"):
        return row["actual_delivery_date"] <= row["scheduled_delivery_date"]
    status = (row.get("delivery_status") or "").casefold().replace("_", " ").strip()
    if any(word in status for word in ("on time", "ontime", "early", "ahead")):
        return True
    if any(word in status for word in ("late", "delay", "overdue", "past due")):
        return False
    return None


def imported_dashboard(dataset_id: str) -> dict[str, Any]:
    metadata, records = read_imported_dataset(dataset_id)
    mappings = metadata["mappings"]
    has_shipments = "shipment_id" in mappings
    has_inventory = "inventory_quantity" in mappings
    has_supplier = bool({"supplier_id", "supplier_name"} & mappings.keys())
    total_orders = len({row["order_id"] for row in records if row.get("order_id")}) if "order_id" in mappings else None

    shipment_rows = [row for row in records if row.get("shipment_id")] if has_shipments else []
    outcomes = [_outcome(row) for row in shipment_rows]
    known_outcomes = [value for value in outcomes if value is not None]
    on_time_rate = round(100 * sum(known_outcomes) / len(known_outcomes), 1) if known_outcomes else None
    delayed_count = sum(value is False for value in known_outcomes) if known_outcomes else None

    supplier_groups: dict[str, list[dict[str, Any]]] = {}
    if has_supplier:
        for row in records:
            key = row.get("supplier_id") or row.get("supplier_name")
            if key:
                supplier_groups.setdefault(str(key), []).append(row)
    suppliers = []
    for key, items in supplier_groups.items():
        sample = items[0]
        supplied_rate = next((row.get("supplier_late_delivery_rate") for row in items if row.get("supplier_late_delivery_rate") is not None), None)
        item_outcomes = [value for value in (_outcome(row) for row in items) if value is not None]
        late_rate = supplied_rate if supplied_rate is not None else (round(100 * sum(not value for value in item_outcomes) / len(item_outcomes), 1) if item_outcomes else None)
        suppliers.append({
            "supplier_id": sample.get("supplier_id") or key,
            "supplier_name": sample.get("supplier_name") or key,
            "late_delivery_rate": late_rate,
            "risk_score": None,
            "risk_level": "NOT SCORED",
            "data_origin": "Imported dataset",
            "synthetic_scenario": None,
        })

    shipment_payload = []
    for row in shipment_rows:
        delay = row.get("delivery_delay_days")
        scheduled = row.get("scheduled_delivery_date")
        actual = row.get("actual_delivery_date")
        if delay is None and actual and scheduled:
            delay = (date.fromisoformat(actual) - date.fromisoformat(scheduled)).days
        ship_date = row.get("ship_date")
        scheduled_days = row.get("scheduled_shipping_days")
        actual_days = row.get("actual_shipping_days")
        if ship_date and scheduled and scheduled_days is None:
            scheduled_days = (date.fromisoformat(scheduled) - date.fromisoformat(ship_date)).days
        if ship_date and actual and actual_days is None:
            actual_days = (date.fromisoformat(actual) - date.fromisoformat(ship_date)).days
        shipment_payload.append({
            "shipment_id": row.get("shipment_id"), "order_id": row.get("order_id"),
            "supplier_count": None, "warehouse_id": row.get("warehouse_id"),
            "order_country": row.get("order_country"), "order_region": row.get("order_region"),
            "ship_datetime": ship_date, "scheduled_shipping_days": scheduled_days,
            "actual_shipping_days": actual_days, "delivery_delay_days": delay,
            "delivery_status": row.get("delivery_status"), "late_delivery_risk": None,
            "shipping_mode": row.get("shipping_mode"), "shipment_status": row.get("delivery_status"),
        })

    trend_records: dict[str, list[bool]] = {}
    if has_shipments:
        for row in shipment_rows:
            date_value = row.get("ship_date") or row.get("actual_delivery_date") or row.get("scheduled_delivery_date")
            result = _outcome(row)
            if date_value and result is not None:
                trend_records.setdefault(date_value[:7], []).append(result)
    trend = [{
        "month": month,
        "shipments": len(values),
        "on_time_rate": round(100 * sum(values) / len(values), 1),
        "delayed": sum(not value for value in values),
    } for month, values in sorted(trend_records.items())[-12:]]

    inventory_rows = []
    if has_inventory:
        for index, row in enumerate(records, start=1):
            inventory_rows.append({
                "inventory_snapshot_id": f"{dataset_id}:{index}",
                "snapshot_date": row.get("snapshot_date") or "",
                "warehouse_id": row.get("warehouse_id") or "",
                "product_id": row.get("product_id") or "",
                "supplier_id": row.get("supplier_id"),
                "estimated_daily_demand": row.get("daily_demand"),
                "current_stock": row.get("inventory_quantity"),
                "reorder_point": row.get("reorder_point"),
                "days_of_cover": None, "supplier_lead_time_days": None,
                "inventory_risk_score": None, "inventory_risk_level": "NOT SCORED",
                "data_origin": "Imported dataset", "synthetic_stock_scenario": None,
                "product_name": row.get("product_name"),
            })

    supplier_count = len(supplier_groups) if has_supplier else None
    available = [FIELD_LABELS[field] for field in mappings]
    unavailable = ["Supplier risk scores", "Inventory risk scores", "Risk event signals"]
    if not has_shipments or not known_outcomes:
        unavailable.append("Delivery performance rate (requires shipment IDs and delivery outcomes)")
    if not has_supplier:
        unavailable.append("Supplier analysis (requires supplier ID or supplier name)")
    if not has_inventory:
        unavailable.append("Inventory analysis (requires inventory quantity)")
    return {
        "data_status": "connected",
        "dataset_kind": "imported",
        "dataset_id": dataset_id,
        "dataset_name": metadata["name"],
        "data_message": f"Imported dataset with {metadata['row_count']:,} validated rows. Only mapped fields are analyzed; unsupported scores and signals are not calculated.",
        "kpis": {
            "total_orders": total_orders,
            "supplier_risk_index": None,
            "supplier_count": supplier_count,
            "high_risk_suppliers": None,
            "shipment_on_time_rate": on_time_rate,
            "total_shipments": len(shipment_rows) if has_shipments else None,
            "delayed_shipments": delayed_count,
            "high_risk_inventory_items": None,
            "inventory_snapshot_rows": len(inventory_rows) if has_inventory else None,
            "inventory_snapshot_date": max((row["snapshot_date"] for row in inventory_rows if row["snapshot_date"]), default=None),
            "high_critical_risk_events": None,
        },
        "trend": trend,
        "suppliers": suppliers,
        "shipments": shipment_payload,
        "inventory": inventory_rows,
        "risk_signals": [],
        "provenance": {
            "source": f"Locally stored import: {metadata['filename']}",
            "mapped_fields": ", ".join(available) or "None",
            "unavailable_analyses": "; ".join(unavailable),
            "risk_scores": "Not calculated for imported data.",
            "forecasting": "Not enabled.",
        },
    }
